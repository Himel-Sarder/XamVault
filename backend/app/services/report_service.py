"""Phase 18 — exported reports (§47, §48).

Both formats are built in memory and streamed; nothing is written to disk, so there are
no temporary files to clean up or to leak between requests.

A report is a *rendering* of data the caller may already read through the API — the same
ownership checks apply, and nothing appears in a report that the caller could not have
fetched as JSON.
"""

from __future__ import annotations

import io
from datetime import UTC, datetime
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfgen.canvas import Canvas
from reportlab.platypus import (
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)
from sqlalchemy.orm import Session

from app.core.exceptions import NotFoundError
from app.models.exam import Exam
from app.models.result import Result
from app.models.user import User
from app.repositories.attempt_repository import AttemptRepository
from app.repositories.exam_repository import ExamRepository
from app.services.app_setting_service import AppSettingService

MARGIN = 18 * mm

#: One palette, so a result sheet and a statement of result look related.
BRAND = colors.HexColor("#1B5E9B")
INK = colors.HexColor("#1B2733")
MUTED = colors.HexColor("#6B7A88")
RULE = colors.HexColor("#D8DEE4")
WASH = colors.HexColor("#F4F7FA")

HEADER_FILL = PatternFill("solid", fgColor="1B5E9B")

_MARKS_NOTE = (
    "Marks shown are as recorded at the time of generation. Where an examination "
    "includes work marked by hand, later moderation may change them."
)
_STATEMENT_NOTE = (
    "This is a computer-generated statement and is valid without a signature. "
    "It reports one examination only and is not a transcript of the programme."
)


def _fmt(value: float | None) -> str:
    if value is None:
        return "-"
    return f"{value:.0f}" if float(value) % 1 == 0 else f"{value:.2f}"


def _safe_filename(text: str) -> str:
    """A filename the browser and the filesystem will both accept."""
    keep = [c if c.isalnum() or c in "-_" else "-" for c in text]
    return "".join(keep).strip("-")[:60] or "report"


def _stage(profile: Any) -> str:
    """"3rd Year, 2nd Semester", or as much of it as is on record."""
    if profile is None:
        return "-"
    parts = []
    if getattr(profile, "study_year", None):
        parts.append(f"{_ordinal(profile.study_year)} Year")
    if getattr(profile, "semester", None):
        parts.append(f"{_ordinal(profile.semester)} Semester")
    return ", ".join(parts) if parts else "-"


def _ordinal(value: int) -> str:
    return {1: "1st", 2: "2nd", 3: "3rd"}.get(value, f"{value}th")


def _styles() -> dict[str, ParagraphStyle]:
    base = getSampleStyleSheet()["Normal"]
    return {
        "doc_title": ParagraphStyle(
            "doc_title",
            parent=base,
            fontName="Helvetica-Bold",
            fontSize=17,
            leading=20,
            textColor=INK,
        ),
        "doc_sub": ParagraphStyle(
            "doc_sub", parent=base, fontSize=10, leading=13, textColor=MUTED
        ),
        "section": ParagraphStyle(
            "section",
            parent=base,
            fontName="Helvetica-Bold",
            fontSize=10,
            textColor=BRAND,
        ),
        "label": ParagraphStyle(
            "label", parent=base, fontSize=7, leading=9, textColor=MUTED
        ),
        "value": ParagraphStyle(
            "value", parent=base, fontSize=9.5, leading=12, textColor=INK
        ),
        "cell": ParagraphStyle("cell", parent=base, fontSize=8, leading=10, textColor=INK),
        "note": ParagraphStyle(
            "note", parent=base, fontSize=7.5, leading=10, textColor=MUTED
        ),
    }


def _facts(pairs: list[tuple[str, str]]) -> Table:
    """The document's metadata as a quiet three-column grid.

    Label above value rather than beside it: a course name and a department name are
    both long, and a fixed label column would either truncate them or waste half the
    page keeping them apart.
    """
    styles = _styles()
    cells = [
        [
            Paragraph(label.upper(), styles["label"]),
            Paragraph(str(value or "-"), styles["value"]),
        ]
        for label, value in pairs
    ]

    columns = 3
    rows: list[list[Any]] = []
    for index in range(0, len(cells), columns):
        chunk = cells[index : index + columns]
        while len(chunk) < columns:
            chunk.append([Paragraph("", styles["label"]), Paragraph("", styles["value"])])
        rows.append([pair[0] for pair in chunk])
        rows.append([pair[1] for pair in chunk])

    width = (A4[0] - 2 * MARGIN) / columns
    table = Table(rows, colWidths=[width] * columns, hAlign="LEFT")
    style = [
        ("LEFTPADDING", (0, 0), (-1, -1), 0),
        ("RIGHTPADDING", (0, 0), (-1, -1), 8),
        ("TOPPADDING", (0, 0), (-1, -1), 0),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 0),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
    ]
    for row in range(0, len(rows), 2):
        style.append(("BOTTOMPADDING", (0, row), (-1, row), 1))
        style.append(("BOTTOMPADDING", (0, row + 1), (-1, row + 1), 8))
    table.setStyle(TableStyle(style))
    return table


def _stat_band(stats: list[tuple[str, str]]) -> Table:
    """The few numbers a reader wants before the table, sized so they are read first."""
    styles = _styles()
    big = ParagraphStyle(
        "stat", parent=styles["value"], fontName="Helvetica-Bold", fontSize=15, leading=18
    )
    labels = [Paragraph(label.upper(), styles["label"]) for label, _ in stats]
    values = [Paragraph(value, big) for _, value in stats]

    width = (A4[0] - 2 * MARGIN) / len(stats)
    table = Table([values, labels], colWidths=[width] * len(stats), hAlign="LEFT")
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, -1), WASH),
                ("LINEABOVE", (0, 0), (-1, 0), 2, BRAND),
                ("TOPPADDING", (0, 0), (-1, 0), 8),
                ("BOTTOMPADDING", (0, 0), (-1, 0), 1),
                ("TOPPADDING", (0, 1), (-1, 1), 0),
                ("BOTTOMPADDING", (0, 1), (-1, 1), 8),
                ("LEFTPADDING", (0, 0), (-1, -1), 10),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ]
        )
    )
    return table


def _score_panel(*, total: str, out_of: str, percentage: str) -> Table:
    """The headline of a statement of result: the mark, in the size it deserves."""
    styles = _styles()
    huge = ParagraphStyle(
        "huge",
        parent=styles["value"],
        fontName="Helvetica-Bold",
        fontSize=30,
        leading=33,
        textColor=BRAND,
    )
    mid = ParagraphStyle(
        "mid",
        parent=styles["value"],
        fontName="Helvetica-Bold",
        fontSize=20,
        leading=24,
        textColor=INK,
    )

    width = (A4[0] - 2 * MARGIN) / 2
    table = Table(
        [
            [
                Paragraph(f"{total} <font size=14>/ {out_of}</font>", huge),
                Paragraph(percentage, mid),
            ],
            [
                Paragraph("MARKS AWARDED", styles["label"]),
                Paragraph("PERCENTAGE", styles["label"]),
            ],
        ],
        colWidths=[width, width],
        hAlign="LEFT",
    )
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, -1), WASH),
                ("LINEABOVE", (0, 0), (-1, 0), 2, BRAND),
                ("LINEBELOW", (0, -1), (-1, -1), 0.3, RULE),
                ("TOPPADDING", (0, 0), (-1, 0), 12),
                ("BOTTOMPADDING", (0, 0), (-1, 0), 0),
                ("TOPPADDING", (0, 1), (-1, 1), 0),
                ("BOTTOMPADDING", (0, 1), (-1, 1), 12),
                ("LEFTPADDING", (0, 0), (-1, -1), 14),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ]
        )
    )
    return table


def _signature_block() -> Table:
    """Two ruled lines. A printed statement ends up on somebody's desk."""
    styles = _styles()
    gutter = 20 * mm
    width = (A4[0] - 2 * MARGIN - gutter) / 2
    # An empty middle column, not merely narrower outer ones: two rules drawn side by
    # side with no gap between them read as a single line across the page.
    table = Table(
        [
            ["", "", ""],
            [
                Paragraph("Controller of Examinations", styles["label"]),
                "",
                Paragraph("Head of Department", styles["label"]),
            ],
        ],
        colWidths=[width, gutter, width],
        hAlign="LEFT",
        rowHeights=[14 * mm, None],
    )
    table.setStyle(
        TableStyle(
            [
                ("LINEBELOW", (0, 0), (0, 0), 0.5, MUTED),
                ("LINEBELOW", (2, 0), (2, 0), 0.5, MUTED),
                ("LEFTPADDING", (0, 0), (-1, -1), 0),
                ("RIGHTPADDING", (0, 0), (-1, -1), 0),
                ("TOPPADDING", (0, 1), (-1, 1), 3),
                ("VALIGN", (0, 0), (-1, -1), "BOTTOM"),
            ]
        )
    )
    return table


def _page_frame(university: str, subtitle: str, kind: str):
    """Draws the band at the top of every page and the rule at the foot.

    Returns a callback rather than a class because ``SimpleDocTemplate`` wants one and
    the three strings never change within a document.
    """

    def draw(canvas: Canvas, doc: Any) -> None:
        width, height = A4
        canvas.saveState()

        # A solid rule across the very top, so a page that has been photocopied or
        # faxed still reads as an official document rather than a screenshot.
        canvas.setFillColor(BRAND)
        canvas.rect(0, height - 6 * mm, width, 6 * mm, stroke=0, fill=1)

        canvas.setFillColor(INK)
        canvas.setFont("Helvetica-Bold", 12)
        canvas.drawString(MARGIN, height - 16 * mm, university)

        if subtitle:
            canvas.setFillColor(MUTED)
            canvas.setFont("Helvetica", 8.5)
            canvas.drawString(MARGIN, height - 21 * mm, subtitle)

        canvas.setFillColor(MUTED)
        canvas.setFont("Helvetica", 8.5)
        canvas.drawRightString(width - MARGIN, height - 16 * mm, kind.upper())

        canvas.setStrokeColor(RULE)
        canvas.setLineWidth(0.5)
        canvas.line(MARGIN, height - 24 * mm, width - MARGIN, height - 24 * mm)

        canvas.line(MARGIN, 16 * mm, width - MARGIN, 16 * mm)
        canvas.setFont("Helvetica", 7.5)
        canvas.setFillColor(MUTED)
        canvas.drawString(
            MARGIN,
            12 * mm,
            f"{university} - generated {datetime.now(UTC).strftime('%d %b %Y %H:%M UTC')}",
        )

        canvas.restoreState()

    return draw


class _NumberedCanvas(Canvas):
    """Adds "Page 2 of 7" to the foot of every page.

    reportlab writes each page as it finishes it, so at that moment the total is not
    known. Pages are held in memory instead and written out at ``save()``, once it is.
    """

    def __init__(self, *args: Any, **kwargs: Any) -> None:
        super().__init__(*args, **kwargs)
        self._saved: list[dict[str, Any]] = []

    def showPage(self) -> None:  # noqa: N802 - reportlab's spelling
        self._saved.append(dict(self.__dict__))
        self._startPage()

    def save(self) -> None:
        total = len(self._saved)
        for state in self._saved:
            self.__dict__.update(state)
            self.setFont("Helvetica", 7.5)
            self.setFillColor(MUTED)
            self.drawRightString(
                A4[0] - MARGIN, 12 * mm, f"Page {self._pageNumber} of {total}"
            )
            super().showPage()
        super().save()


class ReportService:
    def __init__(self, db: Session) -> None:
        self.db = db
        self.exams = ExamRepository(db)
        self.attempts = AttemptRepository(db)

    # --------------------------------------------------------- authorization

    def _owned_exam(self, exam_id: int, actor: User) -> Exam:
        exam = self.exams.get(exam_id)
        if exam is None:
            raise NotFoundError("Exam not found")
        if actor.is_admin or exam.teacher_id == actor.id:
            return exam
        raise NotFoundError("Exam not found")

    def _readable_result(self, result_id: int, actor: User) -> Result:
        result = self.attempts.get_result(result_id)
        if result is None:
            raise NotFoundError("Result not found")

        if actor.is_student:
            # A student may print their own result, and only once it is released —
            # exactly the rule the JSON endpoint applies.
            if result.student_id != actor.id or not result.is_published:
                raise NotFoundError("Result not found")
            return result

        self._owned_exam(result.exam_id, actor)
        return result

    # ------------------------------------------------------------------ rows

    def _rows(self, exam_id: int) -> list[dict[str, Any]]:
        rows: list[dict[str, Any]] = []
        for result in self.attempts.list_results_for_exam(exam_id):
            profile = result.student.student_profile
            rows.append(
                {
                    "student_number": profile.student_id if profile else result.student.user_id,
                    "name": result.student.name,
                    "mcq": float(result.mcq_score),
                    "short": float(result.short_answer_score),
                    "code": float(result.code_score),
                    "total": float(result.total_score),
                    "percentage": float(result.percentage),
                    "status": result.status,
                }
            )
        return rows

    # ------------------------------------------------------------------- PDF

    def _university(self) -> str:
        return AppSettingService(self.db).university_name()

    def _build(
        self,
        buffer: io.BytesIO,
        *,
        pdf_title: str,
        kind: str,
        subtitle: str,
        story: list[Any],
    ) -> None:
        """Puts a story inside the standard frame.

        Every report this service produces shares one page: the same band at the top,
        the same rule and page numbering at the foot. A results sheet and a statement of
        result are different documents, but they should look like they came from the
        same office.
        """
        university = self._university()
        doc = SimpleDocTemplate(
            buffer,
            pagesize=A4,
            title=pdf_title,
            author=university,
            subject=kind,
            leftMargin=MARGIN,
            rightMargin=MARGIN,
            # Room for the band; the story starts below it on every page, not just the
            # first, which is what makes a long table read as one document.
            topMargin=32 * mm,
            bottomMargin=22 * mm,
        )
        frame = _page_frame(university, subtitle, kind)
        doc.build(story, onFirstPage=frame, onLaterPages=frame, canvasmaker=_NumberedCanvas)

    def exam_pdf(self, exam_id: int, actor: User) -> tuple[bytes, str]:
        exam = self._owned_exam(exam_id, actor)
        rows = self._rows(exam_id)
        percentages = [row["percentage"] for row in rows]
        styles = _styles()

        story: list[Any] = [
            Paragraph("Examination Result Sheet", styles["doc_title"]),
            Paragraph(
                f"{exam.course.course_code} &middot; {exam.course.course_name}",
                styles["doc_sub"],
            ),
            Spacer(1, 7 * mm),
            _facts(
                [
                    ("Examination", exam.title),
                    ("Course", f"{exam.course.course_code} - {exam.course.course_name}"),
                    (
                        "Department",
                        exam.course.department.name if exam.course.department else "-",
                    ),
                    ("Examiner", exam.teacher.name),
                    ("Total marks", str(exam.total_marks)),
                    ("Candidates", str(len(rows))),
                ]
            ),
            Spacer(1, 7 * mm),
        ]

        if not rows:
            story.append(
                Paragraph(
                    "No candidate has submitted this examination yet.", styles["note"]
                )
            )
        else:
            average = sum(percentages) / len(percentages)
            passes = sum(1 for value in percentages if value >= 40)
            story += [
                _stat_band(
                    [
                        ("Submitted", str(len(rows))),
                        ("Average", f"{average:.1f}%"),
                        ("Highest", f"{max(percentages):.1f}%"),
                        ("Lowest", f"{min(percentages):.1f}%"),
                        # 40% is the university's own line, not something the platform
                        # decides; it is labelled as a count rather than a verdict.
                        ("At or above 40%", str(passes)),
                    ]
                ),
                Spacer(1, 7 * mm),
                Paragraph("Candidate marks", styles["section"]),
                Spacer(1, 2.5 * mm),
            ]

            head = ["#", "Student ID", "Name", "MCQ", "Short", "Code", "Total", "%", "Status"]
            body = []
            for index, row in enumerate(rows, start=1):
                body.append(
                    [
                        str(index),
                        row["student_number"],
                        Paragraph(row["name"], styles["cell"]),
                        _fmt(row["mcq"]),
                        _fmt(row["short"]),
                        _fmt(row["code"]),
                        f"{_fmt(row['total'])} / {exam.total_marks}",
                        f"{row['percentage']:.1f}",
                        str(row["status"]).replace("_", " ").title(),
                    ]
                )

            table = Table(
                [head, *body],
                repeatRows=1,
                hAlign="LEFT",
                colWidths=[
                    8 * mm,
                    24 * mm,
                    46 * mm,
                    14 * mm,
                    14 * mm,
                    14 * mm,
                    24 * mm,
                    14 * mm,
                    16 * mm,
                ],
            )
            table.setStyle(
                TableStyle(
                    [
                        ("BACKGROUND", (0, 0), (-1, 0), BRAND),
                        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                        ("FONTSIZE", (0, 0), (-1, 0), 7.5),
                        ("FONTSIZE", (0, 1), (-1, -1), 8),
                        ("TEXTCOLOR", (0, 1), (-1, -1), INK),
                        # Numbers right, words left: a column of marks is read by
                        # comparing digits in the same place.
                        ("ALIGN", (3, 0), (-2, -1), "RIGHT"),
                        ("ALIGN", (0, 0), (0, -1), "CENTER"),
                        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, WASH]),
                        # Horizontal rules only. Vertical lines box every cell in and
                        # make a wide table harder to read across, not easier.
                        ("LINEBELOW", (0, 0), (-1, -1), 0.3, RULE),
                        ("LINEBELOW", (0, 0), (-1, 0), 0.8, BRAND),
                        ("TOPPADDING", (0, 0), (-1, -1), 5),
                        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
                        ("LEFTPADDING", (0, 0), (-1, -1), 5),
                        ("RIGHTPADDING", (0, 0), (-1, -1), 5),
                    ]
                )
            )
            story += [table, Spacer(1, 6 * mm), Paragraph(_MARKS_NOTE, styles["note"])]

        buffer = io.BytesIO()
        self._build(
            buffer,
            pdf_title=f"{exam.title} - result sheet",
            kind="Result Sheet",
            subtitle=exam.course.department.name if exam.course.department else "",
            story=story,
        )
        filename = (
            f"{_safe_filename(exam.course.course_code)}-"
            f"{_safe_filename(exam.title)}-results.pdf"
        )
        return buffer.getvalue(), filename

    def result_pdf(self, result_id: int, actor: User) -> tuple[bytes, str]:
        """One student's transcript for a single exam (§47)."""
        result = self._readable_result(result_id, actor)
        exam = result.exam
        profile = result.student.student_profile
        styles = _styles()
        submitted = result.attempt.submitted_at
        percentage = float(result.percentage)

        story: list[Any] = [
            Paragraph("Statement of Result", styles["doc_title"]),
            Paragraph(
                f"{exam.course.course_code} &middot; {exam.title}", styles["doc_sub"]
            ),
            Spacer(1, 7 * mm),
            _facts(
                [
                    ("Candidate", result.student.name),
                    (
                        "Student ID",
                        profile.student_id if profile else result.student.user_id,
                    ),
                    ("Stage", _stage(profile)),
                    ("Course", f"{exam.course.course_code} - {exam.course.course_name}"),
                    (
                        "Department",
                        exam.course.department.name if exam.course.department else "-",
                    ),
                    ("Examination", exam.title),
                    (
                        "Submitted",
                        submitted.strftime("%d %B %Y, %H:%M UTC") if submitted else "-",
                    ),
                ]
            ),
            Spacer(1, 8 * mm),
            _score_panel(
                total=_fmt(float(result.total_score)),
                out_of=str(exam.total_marks),
                percentage=f"{percentage:.1f}%",
            ),
            Spacer(1, 8 * mm),
            Paragraph("Marks by component", styles["section"]),
            Spacer(1, 2.5 * mm),
        ]

        components = Table(
            [
                ["Component", "Marks"],
                ["Multiple choice", _fmt(float(result.mcq_score))],
                ["Short answer", _fmt(float(result.short_answer_score))],
                ["Code", _fmt(float(result.code_score))],
                ["Total", f"{_fmt(float(result.total_score))} / {exam.total_marks}"],
            ],
            colWidths=[110 * mm, 64 * mm],
            hAlign="LEFT",
        )
        components.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), BRAND),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                    ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                    ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"),
                    ("FONTSIZE", (0, 0), (-1, -1), 9),
                    ("TEXTCOLOR", (0, 1), (-1, -1), INK),
                    ("ALIGN", (1, 0), (-1, -1), "RIGHT"),
                    ("ROWBACKGROUNDS", (0, 1), (-1, -2), [colors.white, WASH]),
                    ("BACKGROUND", (0, -1), (-1, -1), WASH),
                    ("LINEBELOW", (0, 0), (-1, -1), 0.3, RULE),
                    ("LINEABOVE", (0, -1), (-1, -1), 0.8, BRAND),
                    ("TOPPADDING", (0, 0), (-1, -1), 6),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
                    ("LEFTPADDING", (0, 0), (-1, -1), 6),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 6),
                ]
            )
        )
        story += [
            components,
            Spacer(1, 10 * mm),
            _signature_block(),
            Spacer(1, 6 * mm),
            Paragraph(_STATEMENT_NOTE, styles["note"]),
        ]

        buffer = io.BytesIO()
        self._build(
            buffer,
            pdf_title=f"{exam.title} - {result.student.name}",
            kind="Statement of Result",
            subtitle=exam.course.department.name if exam.course.department else "",
            story=story,
        )
        filename = (
            f"{_safe_filename(exam.course.course_code)}-"
            f"{_safe_filename(result.student.name)}-result.pdf"
        )
        return buffer.getvalue(), filename


    # ----------------------------------------------------------------- Excel

    def exam_excel(self, exam_id: int, actor: User) -> tuple[bytes, str]:
        """The column set from §48, as a real spreadsheet rather than a CSV renamed."""
        exam = self._owned_exam(exam_id, actor)
        rows = self._rows(exam_id)

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Results"

        sheet["A1"] = self._university()
        sheet["A1"].font = Font(bold=True, size=14)
        sheet["A2"] = f"{exam.course.course_code} - {exam.title}"
        sheet["A3"] = (
            f"Teacher: {exam.teacher.name}    Total marks: {exam.total_marks}    "
            f"Generated: {datetime.now(UTC).strftime('%d %b %Y %H:%M UTC')}"
        )
        sheet["A3"].font = Font(size=9, color="666666")

        headers = [
            "Student ID",
            "Student Name",
            "MCQ Score",
            "Short Answer Score",
            "Code Score",
            "Total Score",
            "Percentage",
            "Status",
        ]
        header_row = 5
        for column, title in enumerate(headers, start=1):
            cell = sheet.cell(row=header_row, column=column, value=title)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = HEADER_FILL
            cell.alignment = Alignment(horizontal="center", vertical="center")

        for offset, row in enumerate(rows, start=header_row + 1):
            # Numbers are written as numbers, not strings: a spreadsheet nobody can sum
            # is not much of an export.
            sheet.cell(row=offset, column=1, value=row["student_number"])
            sheet.cell(row=offset, column=2, value=row["name"])
            sheet.cell(row=offset, column=3, value=row["mcq"])
            sheet.cell(row=offset, column=4, value=row["short"])
            sheet.cell(row=offset, column=5, value=row["code"])
            sheet.cell(row=offset, column=6, value=row["total"])
            percentage = sheet.cell(row=offset, column=7, value=row["percentage"] / 100)
            percentage.number_format = "0.0%"
            sheet.cell(row=offset, column=8, value=row["status"])

        widths = [14, 28, 12, 20, 12, 12, 12, 12]
        for column, width in enumerate(widths, start=1):
            sheet.column_dimensions[get_column_letter(column)].width = width
        # The header stays visible when a long cohort is scrolled.
        sheet.freeze_panes = sheet.cell(row=header_row + 1, column=1)

        buffer = io.BytesIO()
        workbook.save(buffer)
        filename = (
            f"{_safe_filename(exam.course.course_code)}-"
            f"{_safe_filename(exam.title)}-results.xlsx"
        )
        return buffer.getvalue(), filename
