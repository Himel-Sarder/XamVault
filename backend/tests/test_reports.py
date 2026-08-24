"""Phase 18 — PDF and Excel export (§47, §48)."""

from __future__ import annotations

import io
import zipfile

import pytest
from sqlalchemy import select

from app.models.exam import ExamQuestion


@pytest.fixture
def marked_exam(client, api, auth_headers, teacher, student, approved_exam, db):
    """One student has sat and submitted the exam, and the marks are released."""
    headers = auth_headers(student)
    attempt = client.post(
        f"{api}/exams/{approved_exam.id}/start", headers=headers
    ).json()
    link = db.execute(
        select(ExamQuestion).where(ExamQuestion.exam_id == approved_exam.id)
    ).scalars().first()
    correct = next(
        option["id"] for option in link.question_snapshot["options"] if option["is_correct"]
    )
    client.post(
        f"{api}/attempts/{attempt['id']}/answers",
        headers=headers,
        json={
            "exam_question_id": attempt["questions"][0]["exam_question_id"],
            "selected_option_id": correct,
        },
    )
    client.post(f"{api}/attempts/{attempt['id']}/submit", headers=headers)
    client.post(
        f"{api}/exams/{approved_exam.id}/results/publish", headers=auth_headers(teacher)
    )
    result_id = client.get(f"{api}/results", headers=headers).json()[0]["id"]
    return {"exam": approved_exam, "result_id": result_id}


def test_the_exam_pdf_is_a_real_pdf(client, api, auth_headers, teacher, marked_exam):
    response = client.get(
        f"{api}/reports/exams/{marked_exam['exam'].id}/pdf", headers=auth_headers(teacher)
    )

    assert response.status_code == 200
    assert response.headers["content-type"] == "application/pdf"
    assert response.content.startswith(b"%PDF-")
    assert response.content.rstrip().endswith(b"%%EOF")
    assert len(response.content) > 1000


def test_the_pdf_is_offered_as_a_named_download(
    client, api, auth_headers, teacher, marked_exam
):
    response = client.get(
        f"{api}/reports/exams/{marked_exam['exam'].id}/pdf", headers=auth_headers(teacher)
    )
    disposition = response.headers["content-disposition"]

    assert disposition.startswith("attachment;")
    assert disposition.endswith('-results.pdf"')
    # The course code identifies the file without the teacher having to rename it.
    assert marked_exam["exam"].course.course_code.replace("-", "-") in disposition


def test_an_exam_nobody_has_sat_still_produces_a_pdf(
    client, api, auth_headers, teacher, teacher_course, make_published_exam
):
    """An empty result sheet is a valid answer; a 500 is not."""
    exam = make_published_exam(teacher, teacher_course)
    response = client.get(
        f"{api}/reports/exams/{exam.id}/pdf", headers=auth_headers(teacher)
    )

    assert response.status_code == 200
    assert response.content.startswith(b"%PDF-")


def test_the_excel_export_opens_and_holds_the_columns_from_the_spec(
    client, api, auth_headers, teacher, student, marked_exam
):
    from openpyxl import load_workbook

    response = client.get(
        f"{api}/reports/exams/{marked_exam['exam'].id}/excel", headers=auth_headers(teacher)
    )
    assert response.status_code == 200
    assert response.headers["content-type"].endswith("spreadsheetml.sheet")

    # A real xlsx is a zip of XML parts, not a renamed CSV.
    assert zipfile.is_zipfile(io.BytesIO(response.content))

    workbook = load_workbook(io.BytesIO(response.content))
    sheet = workbook["Results"]
    headers = [sheet.cell(row=5, column=c).value for c in range(1, 9)]
    assert headers == [
        "Student ID",
        "Student Name",
        "MCQ Score",
        "Short Answer Score",
        "Code Score",
        "Total Score",
        "Percentage",
        "Status",
    ]

    assert sheet.cell(row=6, column=1).value == student.user_id
    assert sheet.cell(row=6, column=2).value == student.name


def test_the_spreadsheet_holds_numbers_not_text(
    client, api, auth_headers, teacher, marked_exam
):
    """A spreadsheet nobody can sum is not much of an export."""
    from openpyxl import load_workbook

    response = client.get(
        f"{api}/reports/exams/{marked_exam['exam'].id}/excel", headers=auth_headers(teacher)
    )
    sheet = load_workbook(io.BytesIO(response.content))["Results"]

    assert isinstance(sheet.cell(row=6, column=3).value, (int, float)), "mcq score"
    assert isinstance(sheet.cell(row=6, column=6).value, (int, float)), "total"
    percentage = sheet.cell(row=6, column=7)
    assert percentage.value == pytest.approx(1.0), "stored as a fraction"
    assert percentage.number_format == "0.0%", "and formatted as a percentage"


def test_a_student_can_print_their_own_released_result(
    client, api, auth_headers, student, marked_exam
):
    response = client.get(
        f"{api}/reports/results/{marked_exam['result_id']}/pdf", headers=auth_headers(student)
    )
    assert response.status_code == 200
    assert response.content.startswith(b"%PDF-")


def test_a_student_cannot_print_someone_elses_result(
    client, api, auth_headers, other_student, marked_exam
):
    response = client.get(
        f"{api}/reports/results/{marked_exam['result_id']}/pdf",
        headers=auth_headers(other_student),
    )
    assert response.status_code == 404


def test_a_student_cannot_print_an_unreleased_result(
    client, api, auth_headers, teacher, student, approved_exam, db
):
    headers = auth_headers(student)
    attempt = client.post(f"{api}/exams/{approved_exam.id}/start", headers=headers).json()
    client.post(f"{api}/attempts/{attempt['id']}/submit", headers=headers)

    # Not released, so there is nothing for the student to print — the same rule the
    # JSON endpoint applies.
    assert client.get(f"{api}/results", headers=headers).json() == []
    for candidate in range(1, 4):
        assert client.get(
            f"{api}/reports/results/{candidate}/pdf", headers=headers
        ).status_code == 404


def test_the_owning_teacher_can_print_any_of_their_students_results(
    client, api, auth_headers, teacher, marked_exam
):
    response = client.get(
        f"{api}/reports/results/{marked_exam['result_id']}/pdf", headers=auth_headers(teacher)
    )
    assert response.status_code == 200


def test_another_teacher_can_export_nothing(
    client, api, auth_headers, other_teacher, marked_exam
):
    headers = auth_headers(other_teacher)
    exam_id = marked_exam["exam"].id

    assert client.get(f"{api}/reports/exams/{exam_id}/pdf", headers=headers).status_code == 404
    assert client.get(f"{api}/reports/exams/{exam_id}/excel", headers=headers).status_code == 404
    assert client.get(
        f"{api}/reports/results/{marked_exam['result_id']}/pdf", headers=headers
    ).status_code == 404


def test_students_cannot_export_a_whole_cohort(
    client, api, auth_headers, student, marked_exam
):
    headers = auth_headers(student)
    exam_id = marked_exam["exam"].id

    assert client.get(f"{api}/reports/exams/{exam_id}/pdf", headers=headers).status_code == 403
    assert client.get(f"{api}/reports/exams/{exam_id}/excel", headers=headers).status_code == 403


def test_exports_need_authentication(client, api, marked_exam):
    exam_id = marked_exam["exam"].id
    assert client.get(f"{api}/reports/exams/{exam_id}/pdf").status_code == 401
    assert client.get(f"{api}/reports/exams/{exam_id}/excel").status_code == 401
